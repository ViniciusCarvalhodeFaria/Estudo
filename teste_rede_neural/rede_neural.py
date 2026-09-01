import os

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def converter_para_texto(
    caminho_arquivo: str, pasta_saida: str = "arquivos_texto"
) -> str:

    # 1. Verifica se o arquivo existe
    if not os.path.isfile(caminho_arquivo):
        raise FileNotFoundError(
            f"Arquivo não encontrado: {caminho_arquivo}"
        )

    # 2. Descobre a pasta onde este script .py está salvo
    diretorio_script = os.path.dirname(os.path.abspath(__file__))

    # 3. Garante que a pasta de saída seja criada no mesmo diretório do código
    if not os.path.isabs(pasta_saida):
        pasta_saida_completa = os.path.join(diretorio_script, pasta_saida)
    else:
        pasta_saida_completa = pasta_saida

    os.makedirs(pasta_saida_completa, exist_ok=True)

    # 4. Informações do arquivo
    nome_arquivo = os.path.basename(caminho_arquivo)
    nome_sem_extensao, extensao = os.path.splitext(nome_arquivo)
    extensao = extensao.lower()

    # =========================
    # DOCX
    # =========================
    if extensao == ".docx":
        documento = Document(caminho_arquivo)
        textos = [f"ARQUIVO: {nome_arquivo}", "TIPO: DOCX", ""]

        # Extrai os parágrafos
        for paragrafo in documento.paragraphs:
            texto = paragrafo.text.strip()
            if texto:
                textos.append(texto)

        # Extrai tabelas
        for indice, tabela in enumerate(documento.tables, start=1):
            textos.append(f"\n--- TABELA {indice} ---")
            for linha in tabela.rows:
                valores = [celula.text.strip() for celula in linha.cells]
                textos.append(" | ".join(valores))

    # =========================
    # PDF
    # =========================
    elif extensao == ".pdf":
        reader = PdfReader(caminho_arquivo)
        textos = [f"ARQUIVO: {nome_arquivo}", "TIPO: PDF", ""]

        for numero_pagina, pagina in enumerate(reader.pages, start=1):
            texto = pagina.extract_text()
            if texto:
                textos.append(f"\n--- PÁGINA {numero_pagina} ---\n")
                textos.append(texto)

    # =========================
    # EXCEL
    # =========================
    elif extensao in [".xlsx", ".xlsm"]:
        workbook = load_workbook(caminho_arquivo, data_only=True)
        textos = [f"ARQUIVO: {nome_arquivo}", "TIPO: EXCEL", ""]

        for planilha in workbook.worksheets:
            textos.append(f"\n--- PLANILHA: {planilha.title} ---")
            for linha in planilha.iter_rows(values_only=True):
                valores = [
                    str(valor) for valor in linha if valor is not None
                ]
                if valores:
                    textos.append(" | ".join(valores))

    # =========================
    # FORMATO NÃO SUPORTADO
    # =========================
    else:
        raise ValueError(f"Formato não suportado: '{extensao}'")

    # =========================
    # SALVAR ARQUIVO FINAL
    # =========================
    texto_final = "\n".join(textos)

    # Define o caminho absoluto final do arquivo .txt
    caminho_saida = os.path.join(
        pasta_saida_completa, f"{nome_sem_extensao}.txt"
    )

    # Salva o arquivo no disco
    with open(caminho_saida, "w", encoding="utf-8") as arquivo:
        arquivo.write(texto_final)
        print(f"Arquivo convertido com sucesso:\n{caminho_saida}")

    return caminho_saida


# ======================================
# EXEMPLO DE USO
# ======================================
if __name__ == "__main__":
    # Lembre-se de colocar a extensão do arquivo no caminho (ex: .pdf, .docx)
    caminho = "C:/Users/USUARIO/Desktop/Code/teste_rede/curriculo.pdf"

    try:
        arquivo_txt = converter_para_texto(caminho, pasta_saida="saida")
        print(f"\nTexto salvo em: {arquivo_txt}")
    except Exception as erro:
        print(f"Ocorreu um erro: {erro}")